"""
Importa Historial_Partidos_Mitre.xlsx a Supabase.

Uso (desde la raíz del repo):
    python scripts/importar_excel.py --dry-run     valida y muestra qué cambiaría, sin escribir
    python scripts/importar_excel.py               importa y verifica las cifras de control
    python scripts/importar_excel.py --verificar   solo verifica las cifras de control

Lee de .env (o de variables de entorno): SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY y EXCEL_PATH.
La service_role key nunca se imprime ni se sube al repo.

Se puede correr varias veces sin duplicar:
- fuentes, torneos, equipos, jugadores, partidos, posiciones y goleadores de la Liga se
  actualizan por su clave; lo que está en la base y no en el Excel no se borra (solo se avisa).
- goles, incidencias y formaciones de cada partido del Excel quedan exactamente como en el Excel.
- sanciones queda exactamente como en el Excel.
Ignora las columnas calculadas del Excel (las grises) y las hojas LEEME, Resumen, Revisar y Listas.
"""

import argparse
import datetime
import os
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import requests

RAIZ = Path(__file__).resolve().parent.parent

# Mismas listas que los check de sql/01_esquema.sql y sql/04_formaciones.sql
CATEGORIAS = {'Primera', 'Reserva'}
FASES = {'Regular', 'Zona', 'Interzonal', 'Grupos', '32avos', '16avos', 'Octavos', 'Cuartos', 'Semifinal', 'Final',
         'Playoff'}  # Playoff: instancia desconocida
SERIES = {'Ida', 'Vuelta'}
ESTADOS = {'Jugado', 'Programado', 'Suspendido', 'Postergado', 'Ganado por mesa', 'Sin dato'}
COMPLETITUDES = {'Sin dato', 'Inferido', 'Resultado', 'Con goles', 'Completo'}
RESULTADOS = {'G', 'E', 'P'}
TIEMPOS = {'1T', '2T', 'Suplementario', 'Sin dato'}
TIPOS_GOL = {'Jugada', 'Penal', 'Tiro libre', 'En contra', 'Sin dato'}
TIPOS_INCIDENCIA = {'Roja', 'Doble amarilla', 'Amarilla', 'Penal errado', 'Penal atajado',
                    'Expulsión del DT', 'Lesión', 'Partido suspendido', 'Otra'}
ROLES = {'Titular', 'Suplente que entró', 'Suplente sin jugar', 'Convocado (suplente)'}
DIVISIONES = {'Primera', 'Reserva', 'Ambas'}
TIPOS_TORNEO = {'Liga', 'Copa'}

# Cifras de control de docs/PLAN.md (Fase 2)
CONTROL_CAMPANAS = {
    ('2026A', 'Primera'): dict(pj=13, g=2, e=7, p=4, gf=9, gc=10, pts=13),
    ('2026A', 'Reserva'): dict(pj=13, g=5, e=4, p=4, gf=19, gc=17, pts=19),
    ('2026C', 'Primera'): dict(pj=11, g=5, e=1, p=5, gf=14, gc=16, pts=16),
    ('2026C', 'Reserva'): dict(pj=11, g=8, e=2, p=1, gf=24, gc=7, pts=26),
}
CONTROL_GOLEADORES = {
    'Primera': ('Ruiz Díaz', 'Gabriel', 7),
    'Reserva': ('Brisuela Gandolfo', 'Benjamín', 10),
}

# Claves para comparar y actualizar
CLAVES = {
    'fuentes': ('fuente_id',),
    'torneos': ('torneo_id',),
    'equipos': ('equipo_id',),
    'jugadores': ('jugador_id',),
    'partidos': ('partido_id',),
    'posiciones': ('torneo_id', 'categoria', 'tras_fecha'),
    'goleadores_liga': ('torneo_id', 'categoria', 'nombre_en_boletin'),
}
TABLAS_PRINCIPALES = ['fuentes', 'torneos', 'equipos', 'jugadores', 'partidos', 'posiciones', 'goleadores_liga']
TABLAS_POR_PARTIDO = ['goles', 'incidencias', 'formaciones']


class ErrorDeDatos(Exception):
    pass


# ---------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------
def leer_env():
    valores = {}
    archivo = RAIZ / '.env'
    if archivo.exists():
        for linea in archivo.read_text(encoding='utf-8').splitlines():
            linea = linea.strip()
            if not linea or linea.startswith('#') or '=' not in linea:
                continue
            k, v = linea.split('=', 1)
            valores[k.strip()] = v.strip().strip('"').strip("'")
    for k in ('SUPABASE_URL', 'SUPABASE_SERVICE_ROLE_KEY', 'EXCEL_PATH'):
        if os.environ.get(k):
            valores[k] = os.environ[k]
    return valores


# ---------------------------------------------------------------------
# Acceso a Supabase (PostgREST)
# ---------------------------------------------------------------------
class Base:
    LOTE = 500

    def __init__(self, url, clave):
        self.rest = url.rstrip('/') + '/rest/v1/'
        self.sesion = requests.Session()
        self.sesion.headers['apikey'] = clave
        if clave.startswith('eyJ'):  # clave clásica (JWT); las nuevas sb_secret_ van solo en apikey
            self.sesion.headers['Authorization'] = 'Bearer ' + clave

    def _revisar(self, resp, que):
        if resp.status_code >= 300:
            raise SystemExit(f'\nERROR al {que}: {resp.status_code} {resp.text[:500]}')

    def leer(self, tabla, select='*', filtros=None):
        filas, desde = [], 0
        while True:
            params = {'select': select, 'limit': 1000, 'offset': desde, **(filtros or {})}
            resp = self.sesion.get(self.rest + tabla, params=params, timeout=60)
            self._revisar(resp, f'leer {tabla}')
            lote = resp.json()
            filas += lote
            if len(lote) < 1000:
                return filas
            desde += 1000

    def upsert(self, tabla, filas, conflicto):
        for i in range(0, len(filas), self.LOTE):
            resp = self.sesion.post(
                self.rest + tabla, params={'on_conflict': ','.join(conflicto)}, json=filas[i:i + self.LOTE],
                headers={'Prefer': 'resolution=merge-duplicates,return=minimal'}, timeout=120)
            self._revisar(resp, f'guardar en {tabla}')

    def insertar(self, tabla, filas):
        for i in range(0, len(filas), self.LOTE):
            resp = self.sesion.post(self.rest + tabla, json=filas[i:i + self.LOTE],
                                    headers={'Prefer': 'return=minimal'}, timeout=120)
            self._revisar(resp, f'insertar en {tabla}')

    def borrar_por(self, tabla, columna, valores):
        valores = list(valores)
        for i in range(0, len(valores), 50):
            lista = ','.join('"' + str(v).replace('"', '\\"') + '"' for v in valores[i:i + 50])
            resp = self.sesion.delete(self.rest + tabla, params={columna: f'in.({lista})'}, timeout=120)
            self._revisar(resp, f'borrar de {tabla}')

    def borrar_todo(self, tabla, columna_no_nula):
        resp = self.sesion.delete(self.rest + tabla, params={columna_no_nula: 'not.is.null'}, timeout=120)
        self._revisar(resp, f'vaciar {tabla}')


# ---------------------------------------------------------------------
# Lectura del Excel y conversión de valores
# ---------------------------------------------------------------------
class Lector:
    def __init__(self, ruta):
        self.libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        self.errores = []
        self.avisos = []

    def hoja(self, nombre):
        if nombre not in self.libro.sheetnames:
            return []
        filas = self.libro[nombre].iter_rows(values_only=True)
        encabezado = [str(h).strip() if h is not None else None for h in next(filas, [])]
        salida = []
        for nro, valores in enumerate(filas, start=2):
            if not any(v is not None and str(v).strip() != '' for v in valores):
                continue
            fila = dict(zip(encabezado, valores))
            fila['_ubicacion'] = f'{nombre} fila {nro}'
            salida.append(fila)
        return salida

    def error(self, fila, mensaje):
        self.errores.append(f'{fila["_ubicacion"]}: {mensaje}')

    # --- conversores ---
    def texto(self, fila, col, obligatorio=False):
        v = fila.get(col)
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        v = None if v is None else str(v).strip()
        if v == '':
            v = None
        if obligatorio and v is None:
            self.error(fila, f'falta {col}')
        return v

    def entero(self, fila, col, obligatorio=False, minimo=None):
        v = fila.get(col)
        if v is None or (isinstance(v, str) and v.strip() == ''):
            if obligatorio:
                self.error(fila, f'falta {col}')
            return None
        try:
            if isinstance(v, float) and not v.is_integer():
                raise ValueError
            n = int(v)
        except (TypeError, ValueError):
            self.error(fila, f'{col} no es un número entero: {v!r}')
            return None
        if minimo is not None and n < minimo:
            self.error(fila, f'{col} no puede ser menor que {minimo}: {n}')
        return n

    def fecha(self, fila, col):
        v = fila.get(col)
        if v is None or (isinstance(v, str) and v.strip() == ''):
            return None
        if isinstance(v, datetime.datetime):
            return v.date().isoformat()
        if isinstance(v, datetime.date):
            return v.isoformat()
        for formato in ('%d/%m/%Y', '%Y-%m-%d'):
            try:
                return datetime.datetime.strptime(str(v).strip(), formato).date().isoformat()
            except ValueError:
                pass
        self.error(fila, f'{col} no es una fecha: {v!r}')
        return None

    def hora(self, fila, col):
        v = fila.get(col)
        if v is None or (isinstance(v, str) and v.strip() == ''):
            return None
        if isinstance(v, (datetime.time, datetime.datetime)):
            return v.strftime('%H:%M:%S')
        try:
            return datetime.datetime.strptime(str(v).strip(), '%H:%M').strftime('%H:%M:%S')
        except ValueError:
            self.error(fila, f'{col} no es una hora: {v!r}')
            return None

    def lista_ids(self, fila, col):
        v = self.texto(fila, col)
        return [x.strip() for x in v.split(',') if x.strip()] if v else []

    def si_no(self, fila, col):
        v = fila.get(col)
        if v is None or v is False:
            return False
        if v is True:
            return True
        t = str(v).strip().lower()
        if t in ('sí', 'si', 's', 'x', 'true', '1'):
            return True
        if t in ('', 'no', 'n', 'false', '0'):
            return False
        self.error(fila, f'{col} tiene que ser Sí o vacío: {v!r}')
        return False

    def de_lista(self, fila, col, permitidos, obligatorio=False, defecto=None):
        v = self.texto(fila, col) or defecto
        if v is None:
            if obligatorio:
                self.error(fila, f'falta {col}')
            return None
        if v not in permitidos:
            self.error(fila, f'{col} con valor no permitido: {v!r}')
        return v


def leer_excel(ruta, club_propio):
    lx = Lector(ruta)
    d = {}

    d['fuentes'] = [dict(
        fuente_id=lx.texto(r, 'fuente_id', True),
        medio=lx.texto(r, 'medio', True),
        fecha_publicacion=lx.fecha(r, 'fecha_publicacion'),
        que_cubre=lx.texto(r, 'que_cubre'),
        url=lx.texto(r, 'url'),
    ) for r in lx.hoja('Fuentes')]

    d['torneos'] = []
    for r in lx.hoja('Torneos'):
        tid = lx.texto(r, 'torneo_id', True)
        anio = lx.entero(r, 'anio', True)
        orden = lx.entero(r, 'orden')
        if orden is None and tid and anio:
            orden = anio * 10 + {'A': 1, 'C': 2}.get(tid[-1], 5)
        d['torneos'].append(dict(
            torneo_id=tid, anio=anio, nombre=lx.texto(r, 'nombre', True),
            competicion=lx.texto(r, 'competicion') or 'Liga San Martín',
            tipo=lx.de_lista(r, 'tipo', TIPOS_TORNEO, defecto='Liga'),
            organizador=lx.texto(r, 'organizador'),
            formato=lx.texto(r, 'formato'), campeon=lx.texto(r, 'campeon'), notas=lx.texto(r, 'notas'),
            orden=orden, fuente_ids=lx.lista_ids(r, 'fuente_id'),
        ))

    d['equipos'] = []
    for r in lx.hoja('Equipos'):
        eid = lx.texto(r, 'equipo_id', True)
        d['equipos'].append(dict(
            equipo_id=eid, nombre=lx.texto(r, 'nombre', True), nombre_corto=lx.texto(r, 'nombre_corto', True),
            localidad=lx.texto(r, 'localidad'), cancha=lx.texto(r, 'cancha'),
            es_propio=(eid == club_propio), notas=lx.texto(r, 'notas'),
        ))

    d['jugadores'] = [dict(
        jugador_id=lx.texto(r, 'jugador_id', True), apellido=lx.texto(r, 'apellido', True),
        nombre=lx.texto(r, 'nombre'), apodo=lx.texto(r, 'apodo'), posicion=lx.texto(r, 'posicion'),
        division=lx.de_lista(r, 'division', DIVISIONES), fuente_ids=lx.lista_ids(r, 'fuente_id'),
        notas=lx.texto(r, 'notas'),
    ) for r in lx.hoja('Jugadores')]

    # Partidos: se ignoran las columnas calculadas (local, visitante, condicion, rival_id, GF, GC, resultado)
    d['partidos'] = []
    for r in lx.hoja('Partidos'):
        p = dict(
            partido_id=lx.texto(r, 'partido_id', True), torneo_id=lx.texto(r, 'torneo_id', True),
            fase=lx.de_lista(r, 'fase', FASES, defecto='Regular'), zona=lx.texto(r, 'zona'),
            serie=lx.de_lista(r, 'serie', SERIES), fecha_nro=lx.entero(r, 'fecha_nro', minimo=1),
            fecha=lx.fecha(r, 'fecha_partido'), hora=lx.hora(r, 'hora'),
            categoria=lx.de_lista(r, 'categoria', CATEGORIAS, True),
            local_id=lx.texto(r, 'local_id', True), visitante_id=lx.texto(r, 'visitante_id', True),
            goles_local=lx.entero(r, 'goles_local', minimo=0), goles_visitante=lx.entero(r, 'goles_visitante', minimo=0),
            penales_local=lx.entero(r, 'penales_local', minimo=0), penales_visitante=lx.entero(r, 'penales_visitante', minimo=0),
            resultado_informado=lx.de_lista(r, 'resultado_informado', RESULTADOS),
            cancha=lx.texto(r, 'cancha'), arbitro=lx.texto(r, 'arbitro'),
            estado=lx.de_lista(r, 'estado', ESTADOS, defecto='Jugado'),
            completitud=lx.de_lista(r, 'completitud', COMPLETITUDES, defecto='Resultado'),
            fuente_ids=lx.lista_ids(r, 'fuente_id'), observaciones=lx.texto(r, 'observaciones'),
        )
        p['_fila'] = r
        d['partidos'].append(p)

    d['goles'] = []
    for r in lx.hoja('Goles'):
        g = dict(
            gol_id=lx.texto(r, 'gol_id', True), partido_id=lx.texto(r, 'partido_id', True),
            equipo_id=lx.texto(r, 'equipo_id', True), jugador_id=lx.texto(r, 'jugador_id'),
            nombre_en_fuente=lx.texto(r, 'nombre_en_fuente'), minuto=lx.entero(r, 'minuto', minimo=0),
            tiempo=lx.de_lista(r, 'tiempo', TIEMPOS), tipo=lx.de_lista(r, 'tipo', TIPOS_GOL, defecto='Sin dato'),
            observaciones=lx.texto(r, 'observaciones'),
        )
        g['_fila'] = r
        d['goles'].append(g)

    # Incidencias (y la vieja hoja Tarjetas, si tiene filas)
    d['incidencias'] = []
    for hoja in ('Incidencias', 'Tarjetas'):
        for r in lx.hoja(hoja):
            i = dict(
                partido_id=lx.texto(r, 'partido_id', True), equipo_id=lx.texto(r, 'equipo_id'),
                jugador_id=lx.texto(r, 'jugador_id'), nombre_en_fuente=lx.texto(r, 'nombre_en_fuente'),
                tipo=lx.de_lista(r, 'tipo', TIPOS_INCIDENCIA, True) if 'tipo' in r else lx.de_lista(r, 'color', TIPOS_INCIDENCIA, True),
                minuto=lx.entero(r, 'minuto', minimo=0), tiempo=lx.de_lista(r, 'tiempo', TIEMPOS),
                descripcion=lx.texto(r, 'descripcion'),
            )
            i['_fila'] = r
            d['incidencias'].append(i)

    d['formaciones'] = []
    for r in lx.hoja('Formaciones'):
        f = dict(
            partido_id=lx.texto(r, 'partido_id', True), jugador_id=lx.texto(r, 'jugador_id', True),
            rol=lx.de_lista(r, 'rol', ROLES, True), nro_lista=lx.entero(r, 'nro_lista', minimo=1),
            camiseta=lx.entero(r, 'camiseta', minimo=0),
            minuto_entra=lx.entero(r, 'minuto_entra', minimo=0), minuto_sale=lx.entero(r, 'minuto_sale', minimo=0),
            capitan=lx.si_no(r, 'capitan'), observaciones=lx.texto(r, 'observaciones'),
        )
        f['_fila'] = r
        d['formaciones'].append(f)

    d['sanciones'] = []
    for r in lx.hoja('Sanciones'):
        s = dict(
            jugador_id=lx.texto(r, 'jugador_id'), nombre_en_fuente=lx.texto(r, 'nombre_en_fuente'),
            torneo_id=lx.texto(r, 'torneo_id'), fecha_resolucion=lx.fecha(r, 'fecha_resolucion'),
            fechas_suspension=lx.entero(r, 'fechas_suspension', minimo=0), motivo=lx.texto(r, 'motivo'),
            boletin_nro=lx.texto(r, 'boletin_nro'), fuente_ids=lx.lista_ids(r, 'fuente_id'),
        )
        s['_fila'] = r
        d['sanciones'].append(s)

    d['posiciones'] = [dict(
        torneo_id=lx.texto(r, 'torneo_id', True), categoria=lx.de_lista(r, 'categoria', CATEGORIAS, True),
        tras_fecha=lx.entero(r, 'tras_fecha', True, minimo=1), posicion=lx.entero(r, 'posicion', minimo=1),
        pts=lx.entero(r, 'pts'), pj=lx.entero(r, 'PJ', minimo=0), g=lx.entero(r, 'G', minimo=0),
        e=lx.entero(r, 'E', minimo=0), p=lx.entero(r, 'P', minimo=0), dg=lx.entero(r, 'DG'),
        fuente_ids=lx.lista_ids(r, 'fuente_id'), observaciones=lx.texto(r, 'observaciones'),
        _fila=r,
    ) for r in lx.hoja('Posiciones')]

    d['goleadores_liga'] = [dict(
        torneo_id=lx.texto(r, 'torneo_id', True), categoria=lx.de_lista(r, 'categoria', CATEGORIAS, True),
        jugador_id=lx.texto(r, 'jugador_id'), nombre_en_boletin=lx.texto(r, 'nombre_en_boletin', True),
        goles=lx.entero(r, 'goles', True, minimo=0), puesto=lx.entero(r, 'puesto_liga', minimo=1),
        actualizado=lx.fecha(r, 'actualizado'), fuente_ids=lx.lista_ids(r, 'fuente_id'),
        _fila=r,
    ) for r in lx.hoja('Goleadores_Liga')]

    revisar = [r for r in lx.hoja('Revisar') if not lx.texto(r, 'resuelto')]
    if revisar:
        lx.avisos.append(f'La hoja Revisar tiene {len(revisar)} pendientes sin resolver (no se importa).')

    validar(d, lx, club_propio)
    for filas in d.values():
        for f in filas:
            f.pop('_fila', None)
    return d, lx


def validar(d, lx, club_propio):
    def repetidos(tabla, clave):
        c = Counter(tuple(f[k] for k in clave) for f in d[tabla])
        for k, n in c.items():
            if n > 1:
                lx.errores.append(f'{tabla}: {"/".join(map(str, k))} aparece {n} veces')

    for tabla, clave in CLAVES.items():
        repetidos(tabla, clave)
    repetidos('goles', ('gol_id',))
    repetidos('formaciones', ('partido_id', 'jugador_id'))

    fuentes = {f['fuente_id'] for f in d['fuentes']}
    torneos = {t['torneo_id'] for t in d['torneos']}
    equipos = {e['equipo_id'] for e in d['equipos']}
    jugadores = {j['jugador_id'] for j in d['jugadores']}
    partidos = {p['partido_id']: p for p in d['partidos']}

    if club_propio not in equipos:
        lx.errores.append(f'Equipos: no está el club propio ({club_propio})')

    def fuentes_ok(tabla):
        for f in d[tabla]:
            faltan = [x for x in f['fuente_ids'] if x not in fuentes]
            if faltan:
                lx.errores.append(f'{tabla} {f.get(CLAVES.get(tabla, ("",))[0], "")}: fuentes inexistentes {faltan}')

    for tabla in ('torneos', 'jugadores', 'partidos', 'posiciones', 'goleadores_liga', 'sanciones'):
        fuentes_ok(tabla)

    for p in d['partidos']:
        r = p['_fila']
        if p['torneo_id'] and p['torneo_id'] not in torneos:
            lx.error(r, f'torneo inexistente {p["torneo_id"]}')
        for col in ('local_id', 'visitante_id'):
            if p[col] and p[col] not in equipos:
                lx.error(r, f'{col} inexistente {p[col]}')
        if p['local_id'] == p['visitante_id']:
            lx.error(r, 'local y visitante son el mismo equipo')
        if club_propio not in (p['local_id'], p['visitante_id']):
            lx.error(r, f'el club propio ({club_propio}) no juega este partido')
        if (p['goles_local'] is None) != (p['goles_visitante'] is None):
            lx.error(r, 'marcador incompleto (falta uno de los dos goles)')
        if (p['penales_local'] is None) != (p['penales_visitante'] is None):
            lx.error(r, 'penales incompletos')

    for tabla in ('goles', 'incidencias', 'formaciones'):
        for f in d[tabla]:
            r = f['_fila']
            p = partidos.get(f['partido_id'])
            if not p:
                lx.error(r, f'partido inexistente {f["partido_id"]}')
            if f.get('equipo_id'):
                if f['equipo_id'] not in equipos:
                    lx.error(r, f'equipo inexistente {f["equipo_id"]}')
                elif p and f['equipo_id'] not in (p['local_id'], p['visitante_id']):
                    lx.error(r, f'{f["equipo_id"]} no juega el partido {p["partido_id"]}')
            if f.get('jugador_id') and f['jugador_id'] not in jugadores:
                lx.error(r, f'jugador inexistente {f["jugador_id"]}')

    for f in d['sanciones'] + d['posiciones'] + d['goleadores_liga']:
        if f.get('torneo_id') and f['torneo_id'] not in torneos:
            lx.error(f['_fila'], f'torneo inexistente {f["torneo_id"]}')
        if f.get('jugador_id') and f['jugador_id'] not in jugadores:
            lx.error(f['_fila'], f'jugador inexistente {f["jugador_id"]}')

    # Avisos (no frenan la importación)
    propios = Counter(g['partido_id'] for g in d['goles'] if g['equipo_id'] == club_propio)
    for pid, n in propios.items():
        p = partidos.get(pid)
        if p and p['goles_local'] is not None:
            gf = p['goles_local'] if p['local_id'] == club_propio else p['goles_visitante']
            if n != gf:
                lx.avisos.append(f'{pid}: {n} goles propios cargados y el marcador dice {gf}')
    sin_jugador = sum(1 for g in d['goles']
                      if g['equipo_id'] == club_propio and not g['jugador_id'] and g['tipo'] != 'En contra')
    if sin_jugador:
        lx.avisos.append(f'{sin_jugador} goles propios sin jugador asignado.')


# ---------------------------------------------------------------------
# Comparación con la base
# ---------------------------------------------------------------------
def normalizar(v):
    if isinstance(v, list):
        return tuple(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def igual(excel, base):
    return all(normalizar(excel[k]) == normalizar(base.get(k)) for k in excel)


def firma(fila):
    return tuple(sorted((k, normalizar(v)) for k, v in fila.items() if k != 'id'))


def planificar(d, base):
    plan = {}
    for tabla in TABLAS_PRINCIPALES:
        clave = CLAVES[tabla]
        en_base = {tuple(f[k] for k in clave): f for f in base.leer(tabla)}
        nuevas, cambian, iguales = [], [], 0
        for f in d[tabla]:
            k = tuple(f[c] for c in clave)
            if k not in en_base:
                nuevas.append(f)
            elif not igual(f, en_base[k]):
                cambian.append(f)
            else:
                iguales += 1
        en_excel = {tuple(f[c] for c in clave) for f in d[tabla]}
        solo_base = [k for k in en_base if k not in en_excel]
        plan[tabla] = dict(nuevas=nuevas, cambian=cambian, iguales=iguales, solo_base=solo_base)

    ids_partidos = {p['partido_id'] for p in d['partidos']}
    columnas = {t: list(d[t][0].keys()) if d[t] else None for t in TABLAS_POR_PARTIDO}
    for tabla in TABLAS_POR_PARTIDO:
        excel_por_p, base_por_p = defaultdict(list), defaultdict(list)
        for f in d[tabla]:
            excel_por_p[f['partido_id']].append(firma(f))
        cols = columnas[tabla]
        for f in base.leer(tabla):
            if f['partido_id'] in ids_partidos:
                fila = {k: f.get(k) for k in cols} if cols else f
                base_por_p[f['partido_id']].append(firma(fila))
        reemplazar = sorted(pid for pid in set(excel_por_p) | set(base_por_p)
                            if sorted(excel_por_p[pid]) != sorted(base_por_p[pid]))
        filas = [f for f in d[tabla] if f['partido_id'] in set(reemplazar)]
        borradas = sum(len(base_por_p[pid]) for pid in reemplazar)
        plan[tabla] = dict(reemplazar=reemplazar, filas=filas, borradas=borradas,
                           iguales=sum(1 for pid in excel_por_p if pid not in reemplazar))

    cols = list(d['sanciones'][0].keys()) if d['sanciones'] else None
    en_base = base.leer('sanciones')
    firmas_base = sorted(firma({k: f.get(k) for k in cols} if cols else f) for f in en_base)
    cambia = sorted(firma(f) for f in d['sanciones']) != firmas_base
    plan['sanciones'] = dict(reemplazar=cambia, filas=d['sanciones'], borradas=len(en_base) if cambia else 0)

    plan['_borradores'] = [p['partido_id'] for p in base.leer('partidos', 'partido_id', {'publicado': 'eq.false'})]
    return plan


def mostrar_plan(plan):
    print('\nQué cambiaría en la base:')
    print(f'  {"tabla":<17}{"nuevas":>8}{"cambian":>9}{"iguales":>9}   solo en la base (no se borran)')
    for tabla in TABLAS_PRINCIPALES:
        p = plan[tabla]
        solo = len(p['solo_base'])
        extra = f'{solo}: ' + ', '.join('/'.join(map(str, k)) for k in p['solo_base'][:5]) + (' …' if solo > 5 else '') if solo else '-'
        print(f'  {tabla:<17}{len(p["nuevas"]):>8}{len(p["cambian"]):>9}{p["iguales"]:>9}   {extra}')
    print(f'\n  {"por partido":<17}{"partidos a reemplazar":>22}{"filas borradas":>16}{"filas nuevas":>14}')
    for tabla in TABLAS_POR_PARTIDO:
        p = plan[tabla]
        print(f'  {tabla:<17}{len(p["reemplazar"]):>22}{p["borradas"]:>16}{len(p["filas"]):>14}')
    s = plan['sanciones']
    print(f'  {"sanciones":<17}{"se reemplaza" if s["reemplazar"] else "sin cambios":>22}'
          f'{s["borradas"]:>16}{len(s["filas"]) if s["reemplazar"] else 0:>14}')


# ---------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------
def aplicar(plan, base):
    for tabla in TABLAS_PRINCIPALES:
        filas = plan[tabla]['nuevas'] + plan[tabla]['cambian']
        if filas:
            base.upsert(tabla, filas, CLAVES[tabla])
            print(f'  {tabla}: {len(filas)} guardadas')
    for tabla in TABLAS_POR_PARTIDO:
        p = plan[tabla]
        if p['reemplazar']:
            base.borrar_por(tabla, 'partido_id', p['reemplazar'])
            base.insertar(tabla, p['filas'])
            print(f'  {tabla}: {len(p["reemplazar"])} partidos actualizados ({len(p["filas"])} filas)')
    s = plan['sanciones']
    if s['reemplazar']:
        base.borrar_todo('sanciones', 'id')
        base.insertar('sanciones', s['filas'])
        print(f'  sanciones: {len(s["filas"])} filas')


# ---------------------------------------------------------------------
# Cifras de control
# ---------------------------------------------------------------------
def verificar(base):
    print('\nCifras de control (docs/PLAN.md):')
    ok = True
    campanas = {(c['torneo_id'], c['categoria']): c for c in base.leer(
        'v_campanas', 'torneo_id,torneo,categoria,pj,g,e,p,gf,gc,pts',
        {'torneo_id': 'in.(' + ','.join(sorted({k[0] for k in CONTROL_CAMPANAS})) + ')'})}
    for (tid, cat), esperado in CONTROL_CAMPANAS.items():
        real = campanas.get((tid, cat), {})
        dif = {k: (v, real.get(k)) for k, v in esperado.items() if real.get(k) != v}
        nombre = real.get('torneo', tid)
        texto = ' '.join(f'{k.upper()} {real.get(k)}' for k in esperado)
        print(f'  {"OK " if not dif else "MAL"} {nombre} {cat}: {texto}'
              + (f'   esperado: {dif}' if dif else ''))
        ok &= not dif

    octavos = base.leer('v_partidos', 'local,visitante,goles_local,goles_visitante,penales_local,penales_visitante',
                        {'torneo_id': 'eq.2026A', 'fase': 'eq.Octavos', 'categoria': 'eq.Primera'})
    o = octavos[0] if octavos else {}
    bien = (o.get('local') == 'General' and (o.get('goles_local'), o.get('goles_visitante')) == (2, 2)
            and (o.get('penales_local'), o.get('penales_visitante')) == (5, 3))
    print(f'  {"OK " if bien else "MAL"} Octavos Apertura 2026: {o.get("local")} {o.get("goles_local")}-'
          f'{o.get("goles_visitante")} {o.get("visitante")}, penales {o.get("penales_local")}-{o.get("penales_visitante")}')
    ok &= bien

    goles = base.leer('v_goleadores', 'jugador_id,apellido,nombre,categoria,goles', {'anio': 'eq.2026'})
    for cat, (apellido, nombre, esperado) in CONTROL_GOLEADORES.items():
        suma = Counter()
        datos = {}
        for g in goles:
            if g['categoria'] == cat:
                suma[g['jugador_id']] += g['goles']
                datos[g['jugador_id']] = g
        jid, n = suma.most_common(1)[0] if suma else (None, 0)
        j = datos.get(jid, {})
        bien = j.get('apellido') == apellido and (j.get('nombre') or '').startswith(nombre) and n == esperado
        print(f'  {"OK " if bien else "MAL"} Goleador {cat} 2026: {j.get("apellido")}, {j.get("nombre")} ({n} goles)')
        ok &= bien
    return ok


# ---------------------------------------------------------------------
def main():
    sys.stdout.reconfigure(encoding='utf-8')
    ap = argparse.ArgumentParser(description='Importa el Excel del historial a Supabase.')
    ap.add_argument('--dry-run', action='store_true', help='valida y muestra qué cambiaría, sin escribir')
    ap.add_argument('--verificar', action='store_true', help='solo verifica las cifras de control')
    ap.add_argument('--excel', help='ruta del Excel (por defecto, EXCEL_PATH de .env)')
    ap.add_argument('--club-propio', default='MIT', help='equipo_id del club propio (por defecto MIT)')
    ap.add_argument('--forzar', action='store_true',
                    help='importar aunque la base tenga partidos que no están en el Excel o borradores')
    args = ap.parse_args()

    env = leer_env()
    faltan = [k for k in ('SUPABASE_URL', 'SUPABASE_SERVICE_ROLE_KEY') if not env.get(k)]
    if faltan:
        raise SystemExit(f'Faltan en .env: {", ".join(faltan)} (ver .env.example)')
    base = Base(env['SUPABASE_URL'], env['SUPABASE_SERVICE_ROLE_KEY'])

    if args.verificar:
        sys.exit(0 if verificar(base) else 1)

    ruta = args.excel or env.get('EXCEL_PATH')
    if not ruta or not Path(ruta).exists():
        raise SystemExit(f'No encuentro el Excel: {ruta!r} (EXCEL_PATH en .env o --excel)')
    print(f'Leyendo {ruta}')
    d, lx = leer_excel(ruta, args.club_propio)
    print('Filas en el Excel: ' + ', '.join(f'{t} {len(f)}' for t, f in d.items()))

    for a in lx.avisos:
        print('  aviso:', a)
    if lx.errores:
        print(f'\n{len(lx.errores)} errores. No se importa nada:')
        for e in lx.errores:
            print('  -', e)
        sys.exit(1)
    print('Validación OK.')

    plan = planificar(d, base)
    mostrar_plan(plan)

    frenos = []
    if plan['partidos']['solo_base']:
        frenos.append(f'{len(plan["partidos"]["solo_base"])} partidos están en la base y no en el Excel '
                      '(¿cargados desde la app?)')
    if plan['_borradores']:
        frenos.append(f'{len(plan["_borradores"])} partidos en borrador en la base (cargados desde la app)')
    if frenos:
        print('\nATENCIÓN: la base parece tener datos cargados desde la app. Importar podría pisarlos.')
        for f in frenos:
            print('  -', f)
        if not args.dry_run and not args.forzar:
            raise SystemExit('No se importa. Si estás seguro, volvé a correr con --forzar.')

    if args.dry_run:
        print('\nDry-run: no se escribió nada.')
        return

    print('\nImportando…')
    aplicar(plan, base)
    print('Importación terminada.')
    sys.exit(0 if verificar(base) else 1)


if __name__ == '__main__':
    main()
